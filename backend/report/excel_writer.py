
import openpyxl
from openpyxl.chart import LineChart, Reference

def write_efficiency_report(df, xlsx_path):
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Data"
    ws.append(list(df.columns))
    for _,r in df.iterrows(): ws.append(list(r.values))

    ws2=wb.create_sheet("Eff_By_Combo")
    combos=sorted(set((r["i_ch1"], r["i23"]) for _, r in df.iterrows()))
    ws2.append(["vin"]+[f"i1={i1} i23={i23}" for (i1,i23) in combos])
    vins=sorted(df["vin"].unique())
    for v in vins:
        row=[v]
        for (i1,i23) in combos:
            m=df[(df["vin"]==v)&(df["i_ch1"]==i1)&(df["i23"]==i23)]
            row.append(m["efficiency"].iloc[0] if len(m)>0 else None)
        ws2.append(row)
    chart=LineChart(); chart.title="Efficiency vs Vin (by loads)"; chart.y_axis.title="Efficiency"; chart.x_axis.title="Vin (V)"
    data=Reference(ws2,min_col=2,min_row=1,max_col=1+len(combos),max_row=1+len(vins))
    cats=Reference(ws2,min_col=1,min_row=2,max_row=1+len(vins))
    chart.add_data(data,titles_from_data=True); chart.set_categories(cats); ws2.add_chart(chart,"H2")
    wb.save(xlsx_path)
